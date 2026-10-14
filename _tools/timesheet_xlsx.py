#!/usr/bin/env python3
"""
Builds the downloadable Excel timesheet templates (needs openpyxl):
    python _tools/timesheet_xlsx.py
Writes assets/files/weekly-timesheet.xlsx and assets/files/biweekly-timesheet.xlsx.
Each workbook has a blank "Timesheet" sheet and a filled-in "Example" sheet.
"""
import datetime
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, 'assets', 'files')

BLUE = '1D4ED8'
HEAD = PatternFill('solid', fgColor='EFF6FF')
TOTAL = PatternFill('solid', fgColor='F1F5F9')
THIN = Side(style='thin', color='CBD5E1')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
DAYS = ['Monday', 'Tuesday', 'Wednesday', 'Thursday', 'Friday', 'Saturday', 'Sunday']
TIME_FMT = 'h:mm AM/PM'


def hours_formula(r):
    # Out − In (crossing midnight with MOD) minus the lunch break, in decimal hours.
    return (f'=IF(OR(C{r}="",F{r}=""),0,ROUND((MOD(F{r}-C{r},1)'
            f'-IF(OR(D{r}="",E{r}=""),0,MOD(E{r}-D{r},1)))*24,2))')


def sheet(wb, title, weeks, example):
    ws = wb.create_sheet(title)
    ws.sheet_view.showGridLines = False
    widths = [14, 13, 13, 13, 13, 13, 14, 3, 26, 14]
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    ws['A1'] = 'Weekly Timesheet' if weeks == 1 else 'Biweekly Timesheet'
    ws['A1'].font = Font(size=18, bold=True, color=BLUE)
    ws['A2'] = 'Free template from workhourmath.com — times in 12-hour format, e.g. 9:00 AM'
    ws['A2'].font = Font(size=9, italic=True, color='5B6577')
    labels = [('A4', 'Employee'), ('A5', 'Week starting'), ('E4', 'Hourly rate ($)'), ('E5', 'Overtime after (hrs/week)')]
    for cell, text in labels:
        ws[cell] = text
        ws[cell].font = Font(bold=True)
    ws['G5'] = 40
    if example:
        ws['B4'] = 'Example employee'
        ws['G4'] = 18.5
        ws['B5'] = datetime.date(2026, 10, 19)
    ws['B5'].number_format = 'mm/dd/yyyy'
    ws['G4'].number_format = '$#,##0.00'
    for c in ('B4', 'B5', 'G4', 'G5'):
        ws[c].border = BOX

    head = ['Day', 'Date', 'Time in', 'Lunch out', 'Lunch in', 'Time out', 'Hours']
    tv = DataValidation(type='time', operator='between', formula1='0', formula2='0.999988426',
                        error='Enter a time such as 9:00 AM', errorTitle='Time', allow_blank=True)
    ws.add_data_validation(tv)
    row = 7
    week_totals = []
    sample = [('8:45 AM', '12:30 PM', '1:00 PM', '5:15 PM')] * 5 + [('8:00 AM', '', '', '12:30 PM'), ('', '', '', '')]
    for w in range(weeks):
        if weeks > 1:
            ws.cell(row=row, column=1, value=f'Week {w + 1}').font = Font(bold=True, color=BLUE)
            row += 1
        for i, h in enumerate(head):
            c = ws.cell(row=row, column=i + 1, value=h)
            c.font = Font(bold=True)
            c.fill = HEAD
            c.border = BOX
            c.alignment = Alignment(horizontal='center')
        row += 1
        first = row
        for d in range(7):
            ws.cell(row=row, column=1, value=DAYS[d]).border = BOX
            dc = ws.cell(row=row, column=2, value=f'=IF($B$5="","",$B$5+{w * 7 + d})')
            dc.number_format = 'mm/dd'
            dc.border = BOX
            for col in range(3, 7):
                c = ws.cell(row=row, column=col)
                c.number_format = TIME_FMT
                c.border = BOX
                tv.add(c)
                if example and sample[d][col - 3]:
                    t = datetime.datetime.strptime(sample[d][col - 3], '%I:%M %p').time()
                    c.value = t
            hc = ws.cell(row=row, column=7, value=hours_formula(row))
            hc.number_format = '0.00'
            hc.border = BOX
            row += 1
        tot = ws.cell(row=row, column=6, value='Week total' if weeks > 1 else 'Total hours')
        tot.font = Font(bold=True)
        t = ws.cell(row=row, column=7, value=f'=SUM(G{first}:G{row - 1})')
        t.number_format = '0.00'
        t.font = Font(bold=True)
        t.fill = TOTAL
        t.border = BOX
        week_totals.append(row)
        row += 2

    # Summary: overtime is counted per workweek (FLSA), then summed.
    summary = [
        ('Total hours', '=' + '+'.join(f'G{r}' for r in week_totals), '0.00'),
        ('Regular hours', '=' + '+'.join(f'MIN(G{r},$G$5)' for r in week_totals), '0.00'),
        ('Overtime hours (1.5×)', '=' + '+'.join(f'MAX(0,G{r}-$G$5)' for r in week_totals), '0.00'),
        ('Regular pay', None, '$#,##0.00'),
        ('Overtime pay', None, '$#,##0.00'),
        ('Gross pay', None, '$#,##0.00'),
    ]
    start = row
    for i, (label, formula, fmt) in enumerate(summary):
        r = start + i
        ws.cell(row=r, column=9, value=label).font = Font(bold=(label in ('Total hours', 'Gross pay')))
        if label == 'Regular pay':
            formula = f'=IF($G$4="","",J{start + 1}*$G$4)'
        elif label == 'Overtime pay':
            formula = f'=IF($G$4="","",J{start + 2}*$G$4*1.5)'
        elif label == 'Gross pay':
            formula = f'=IF($G$4="","",J{start + 3}+J{start + 4})'
        c = ws.cell(row=r, column=10, value=formula)
        c.number_format = fmt
        c.border = BOX
        if label in ('Total hours', 'Gross pay'):
            c.font = Font(bold=True)
            c.fill = TOTAL
    ws.cell(row=start + 7, column=9, value='Gross pay is before taxes. Overtime here follows the federal 40-hour weekly rule;').font = Font(size=9, color='5B6577')
    ws.cell(row=start + 8, column=9, value='some states also have daily overtime — see workhourmath.com/guides/overtime-rules-by-state.html').font = Font(size=9, color='5B6577')
    ws.cell(row=start + 10, column=1, value='Supervisor signature: ______________________     Employee signature: ______________________').font = Font(size=10)

    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = f'A1:J{start + 11}'
    return ws


def build(weeks, name):
    wb = Workbook()
    wb.remove(wb.active)
    sheet(wb, 'Timesheet', weeks, example=False)
    sheet(wb, 'Example', weeks, example=True)
    wb.properties.title = name
    wb.properties.creator = 'WorkHourMath'
    os.makedirs(OUT, exist_ok=True)
    path = os.path.join(OUT, name)
    wb.save(path)
    print('wrote', os.path.relpath(path, ROOT))


if __name__ == '__main__':
    build(1, 'weekly-timesheet.xlsx')
    build(2, 'biweekly-timesheet.xlsx')
